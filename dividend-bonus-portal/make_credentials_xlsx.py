"""dist/credentials.csv -> dist/Dividend_Bonus_Portal_Logins.xlsx (hand out; never publish or commit)."""
import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HERE = Path(__file__).resolve().parent
URL = 'https://nishanthm-svg.github.io/shreeja-lms/dividend-bonus/'
rows = list(csv.reader(open(HERE / 'dist' / 'credentials.csv', encoding='utf-8-sig')))
wb = Workbook(); ws = wb.active; ws.title = 'Logins'
ws.append(['Members Dividend & Bonus portal: ' + URL])
ws['A1'].font = Font(bold=True)
ws.append(['Type the Login ID column (e.g. ADMIN, ACO05) in the Login ID box, not the name. Passwords are lowercase.'])
ws.append([])
for r in rows:
    ws.append([int(v) if v.isdigit() and i >= 4 else v for i, v in enumerate(r)])
for c in ws[4]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='047857')
for col, w in zip('ABCDEF', (11, 8, 26, 12, 8, 10)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'A5'
ws.auto_filter.ref = f'A4:F{ws.max_row}'
wb.save(HERE / 'dist' / 'Dividend_Bonus_Portal_Logins.xlsx')
print('dist/Dividend_Bonus_Portal_Logins.xlsx')
